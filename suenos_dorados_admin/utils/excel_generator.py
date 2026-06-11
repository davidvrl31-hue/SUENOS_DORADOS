from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


def export_inventory_to_excel(rows, output_dir: Path) -> Path:
    """Genera un Excel liviano de inventario desde filas ya consultadas."""
    output_dir.mkdir(parents=True, exist_ok=True)
    file_path = output_dir / f"inventario_{datetime.now():%Y%m%d_%H%M%S}.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Inventario"

    headers = [
        "ID variante",
        "Producto",
        "Categoria",
        "Medida",
        "Color",
        "SKU",
        "Referencia",
        "Precio",
        "Stock",
        "Estado",
    ]
    ws.append(headers)

    for row in rows:
        ws.append(
            [
                row.get("id_variante"),
                row.get("producto"),
                row.get("categoria"),
                row.get("medida"),
                row.get("color"),
                row.get("sku"),
                row.get("referencia"),
                float(row.get("precio") or 0),
                int(row.get("stock") or 0),
                "Activo" if row.get("estado") else "Inactivo",
            ]
        )

    _format_inventory_sheet(ws, headers)
    wb.save(file_path)
    return file_path


def _format_inventory_sheet(ws, headers: list[str]) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    header_fill = PatternFill("solid", fgColor="1F2937")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    widths = {
        "A": 12,
        "B": 28,
        "C": 18,
        "D": 18,
        "E": 18,
        "F": 34,
        "G": 18,
        "H": 14,
        "I": 10,
        "J": 12,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    for row in ws.iter_rows(min_row=2, max_col=len(headers)):
        for cell in row:
            cell.alignment = Alignment(vertical="center")
        row[7].number_format = '"$"#,##0.00'


